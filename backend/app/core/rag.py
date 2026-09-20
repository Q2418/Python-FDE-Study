import io
import re
from pathlib import Path

from ..config import (
    PROJECT_ALLOWED_EXTENSIONS,
    PROJECT_CHUNK_OVERLAP,
    PROJECT_CHUNK_SIZE,
    PROJECT_ROOT,
    PROJECT_SCAN_DIRS,
    PROJECT_SCAN_FILES,
    RAG_CHUNK_OVERLAP,
    RAG_CHUNK_SIZE,
    RAG_MIN_SCORE,
    RAG_TOP_K,
)


def parse_document(filename: str, data: bytes) -> str:
    """解析知识库文档：支持 txt / md / docx"""
    lower = filename.lower()
    if lower.endswith(".docx"):
        from docx import Document

        document = Document(io.BytesIO(data))
        return "\n".join(p.text for p in document.paragraphs if p.text.strip())
    return data.decode("utf-8", errors="ignore")


def split_text(text: str, chunk_size: int = None, overlap: int = None) -> list:
    """按固定长度切片，带重叠避免上下文断裂（参数可配置调优）"""
    chunk_size = chunk_size or RAG_CHUNK_SIZE
    overlap = RAG_CHUNK_OVERLAP if overlap is None else overlap
    normalized = re.sub(r"\s+", " ", text).strip()
    if not normalized:
        return []
    chunks = []
    start = 0
    while start < len(normalized):
        end = start + chunk_size
        chunks.append(normalized[start:end])
        if end >= len(normalized):
            break
        start = end - overlap
    return chunks


def _bigrams(text: str) -> set:
    cleaned = re.sub(r"[^\w\u4e00-\u9fff]", "", text.lower())
    if len(cleaned) < 2:
        return {cleaned} if cleaned else set()
    return {cleaned[i:i + 2] for i in range(len(cleaned) - 1)}


def score_chunk(query: str, chunk: str) -> float:
    """简易相似度：查询与切片的字符二元组重合率（无需外部向量库）"""
    query_grams = _bigrams(query)
    if not query_grams:
        return 0.0
    chunk_grams = _bigrams(chunk)
    if not chunk_grams:
        return 0.0
    return len(query_grams & chunk_grams) / len(query_grams)


def retrieve(db, query: str, top_k: int = None, min_score: float = None) -> list:
    from ..models.ai import AiChunk, AiDocument

    top_k = top_k or RAG_TOP_K
    min_score = RAG_MIN_SCORE if min_score is None else min_score
    rows = (
        db.query(AiChunk, AiDocument.filename)
        .join(AiDocument, AiChunk.document_id == AiDocument.id)
        .all()
    )
    scored = [(score_chunk(query, chunk.content), chunk, filename) for chunk, filename in rows]
    scored.sort(key=lambda item: item[0], reverse=True)
    return [item for item in scored[:top_k] if item[0] >= min_score]


def collect_project_files() -> list:
    """收集需要入库的项目资料：表结构 SQL、README、docs、后端代码"""
    files = []
    for relative in PROJECT_SCAN_FILES:
        path = PROJECT_ROOT / relative
        if path.is_file():
            files.append(path)
    for folder in PROJECT_SCAN_DIRS:
        base = PROJECT_ROOT / folder
        if not base.is_dir():
            continue
        for path in sorted(base.rglob("*")):
            if not path.is_file():
                continue
            if "__pycache__" in path.parts:
                continue
            if path.suffix.lower() in PROJECT_ALLOWED_EXTENSIONS:
                files.append(path)
    return files


def reindex_project(db, uploader_id: int = None) -> dict:
    """重建项目知识库：清空旧的项目切片，重新扫描入库"""
    from ..models.ai import AiChunk, AiDocument

    old_docs = db.query(AiDocument).filter(AiDocument.category == "project").all()
    old_ids = [doc.id for doc in old_docs]
    if old_ids:
        db.query(AiChunk).filter(AiChunk.document_id.in_(old_ids)).delete(synchronize_session=False)
        for doc in old_docs:
            db.delete(doc)
        db.commit()

    indexed_files = 0
    total_chunks = 0
    for path in collect_project_files():
        try:
            content = path.read_text(encoding="utf-8", errors="ignore")
        except OSError:
            continue
        chunks = split_text(content, chunk_size=PROJECT_CHUNK_SIZE, overlap=PROJECT_CHUNK_OVERLAP)
        if not chunks:
            continue
        relative = str(Path(path).relative_to(PROJECT_ROOT)).replace("\\", "/")
        document = AiDocument(
            filename=relative,
            file_path=str(path),
            file_size=path.stat().st_size,
            chunk_count=len(chunks),
            category="project",
            uploader_id=uploader_id,
        )
        db.add(document)
        db.flush()
        db.add_all([AiChunk(document_id=document.id, seq=index, content=chunk) for index, chunk in enumerate(chunks)])
        indexed_files += 1
        total_chunks += len(chunks)
    db.commit()
    return {"files": indexed_files, "chunks": total_chunks}
